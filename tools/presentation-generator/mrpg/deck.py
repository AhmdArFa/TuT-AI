"""PPTX rendering: turn an :class:`Outline` into a polished 16:9 deck.

The default theme follows the TuT visual identity: deep navy background,
royal-gold accent, clean sans typography. Rendering is pure python-pptx,
so decks are native, editable PowerPoint files.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
NAVY_LIGHT = RGBColor(0x12, 0x33, 0x5E)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
GOLD_LIGHT = RGBColor(0xE8, 0xC9, 0x6A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xD9, 0xE2, 0xEC)
FONT = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _fill_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _rect(slide, x, y, w, h, color: RGBColor):
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def _textbox(slide, x, y, w, h):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    return box, frame


def _set_run(run, text: str, size: int, color: RGBColor, bold: bool = False) -> None:
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = FONT


def _title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = _blank(prs)
    _fill_background(slide, NAVY)
    _rect(slide, Inches(0), Inches(3.05), SLIDE_W, Pt(3), GOLD)
    _, frame = _textbox(slide, Inches(0.9), Inches(1.55), Inches(11.5), Inches(1.5))
    _set_run(frame.paragraphs[0].add_run(), title, 40, WHITE, bold=True)
    if subtitle:
        _, sub = _textbox(slide, Inches(0.9), Inches(3.35), Inches(11.5), Inches(1.0))
        _set_run(sub.paragraphs[0].add_run(), subtitle, 20, GOLD_LIGHT)
    _, brand = _textbox(slide, Inches(0.9), Inches(6.6), Inches(11.5), Inches(0.5))
    _set_run(
        brand.paragraphs[0].add_run(),
        "Multi-Resource Presentation Generator",
        12,
        OFF_WHITE,
    )


def _content_slide(prs: Presentation, slide_data, number: int, total: int) -> None:
    slide = _blank(prs)
    _fill_background(slide, NAVY)
    _rect(slide, Inches(0), Inches(0), Inches(0.14), SLIDE_H, GOLD)

    _, title_frame = _textbox(
        slide, Inches(0.55), Inches(0.35), Inches(12.2), Inches(0.95)
    )
    _set_run(
        title_frame.paragraphs[0].add_run(), slide_data.title, 28, WHITE, bold=True
    )
    _rect(slide, Inches(0.57), Inches(1.22), Inches(1.6), Pt(3), GOLD)

    _, frame = _textbox(slide, Inches(0.6), Inches(1.55), Inches(12.1), Inches(4.9))
    for index, bullet in enumerate(slide_data.bullets):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.space_after = Pt(14)
        marker = para.add_run()
        _set_run(marker, "▸ ", 18, GOLD, bold=True)
        text_run = para.add_run()
        _set_run(text_run, bullet, 18, OFF_WHITE)

    footer_parts = []
    if slide_data.sources:
        footer_parts.append("Sources: " + ", ".join(slide_data.sources))
    footer_parts.append(f"{number} / {total}")
    _, footer = _textbox(slide, Inches(0.6), Inches(6.95), Inches(12.1), Inches(0.4))
    para = footer.paragraphs[0]
    para.alignment = PP_ALIGN.RIGHT if not slide_data.sources else PP_ALIGN.LEFT
    _set_run(
        para.add_run(), "    ·    ".join(footer_parts), 10, RGBColor(0x8A, 0x9B, 0xB5)
    )

    if slide_data.notes:
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = slide_data.notes


def _closing_slide(prs: Presentation, sources: list[str]) -> None:
    slide = _blank(prs)
    _fill_background(slide, NAVY)
    _rect(slide, Inches(0), Inches(2.6), SLIDE_W, Pt(3), GOLD)
    _, frame = _textbox(slide, Inches(0.9), Inches(1.5), Inches(11.5), Inches(1.0))
    _set_run(frame.paragraphs[0].add_run(), "Thank You", 36, WHITE, bold=True)
    if sources:
        _, frame = _textbox(slide, Inches(0.9), Inches(2.95), Inches(11.5), Inches(3.5))
        head = frame.paragraphs[0].add_run()
        _set_run(head, "Resources used", 16, GOLD_LIGHT, bold=True)
        for source in sources:
            para = frame.add_paragraph()
            para.space_before = Pt(6)
            _set_run(para.add_run(), f"• {source}", 13, OFF_WHITE)


def render_deck(outline, output_path: str | Path) -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    _title_slide(prs, outline.title, outline.subtitle)
    total = len(outline.slides)
    for index, slide_data in enumerate(outline.slides, start=1):
        _content_slide(prs, slide_data, index, total)

    all_sources: list[str] = []
    for slide_data in outline.slides:
        for source in slide_data.sources:
            if source not in all_sources:
                all_sources.append(source)
    _closing_slide(prs, all_sources)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return output
