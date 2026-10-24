"""Resource ingestion for the Multi-Resource Presentation Generator.

Extracts clean text from local files (offline resources) and web pages
(online resources). Every extractor returns a :class:`Resource` with the
source origin preserved, so generated slides can cite where content
came from.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md", ".markdown"}


@dataclass
class Resource:
    """One ingested resource: origin metadata plus extracted text blocks."""

    source: str  # file path or URL
    kind: str  # pdf | docx | pptx | text | web
    title: str = ""
    blocks: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(b for b in self.blocks if b.strip())

    @property
    def origin_label(self) -> str:
        if self.kind == "web":
            return self.source
        return Path(self.source).name


def _clean(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text or "")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pretty_stem(path: Path) -> str:
    words = path.stem.replace("_", " ").replace("-", " ").split()
    return " ".join(w.capitalize() if w.islower() else w for w in words)


def _from_pdf(path: Path) -> Resource:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks = []
    for page in reader.pages:
        page_text = _clean(page.extract_text() or "")
        if page_text:
            blocks.append(page_text)
    title = ""
    if reader.metadata and reader.metadata.title:
        title = str(reader.metadata.title)
    return Resource(
        source=str(path), kind="pdf", title=title or _pretty_stem(path), blocks=blocks
    )


def _from_docx(path: Path) -> Resource:
    import docx

    document = docx.Document(str(path))
    blocks = []
    for para in document.paragraphs:
        line = _clean(para.text)
        if line:
            blocks.append(line)
    for table in document.tables:
        for row in table.rows:
            cells = [_clean(c.text) for c in row.cells if _clean(c.text)]
            if cells:
                blocks.append(" | ".join(cells))
    return Resource(
        source=str(path), kind="docx", title=_pretty_stem(path), blocks=blocks
    )


def _from_pptx(path: Path) -> Resource:
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks = []
    for index, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = _clean("".join(run.text for run in para.runs))
                    if line:
                        parts.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [_clean(c.text) for c in row.cells if _clean(c.text)]
                    if cells:
                        parts.append(" | ".join(cells))
        if parts:
            blocks.append(f"[Slide {index}] " + "\n".join(parts))
    return Resource(
        source=str(path), kind="pptx", title=_pretty_stem(path), blocks=blocks
    )


def _from_text(path: Path) -> Resource:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    blocks = [b for b in (_clean(p) for p in re.split(r"\n\s*\n", raw)) if b]
    title = _pretty_stem(path)
    if blocks and blocks[0].startswith("#"):
        title = blocks[0].lstrip("# ").strip()
    return Resource(source=str(path), kind="text", title=title, blocks=blocks)


def _from_web(url: str) -> Resource:
    import requests
    from bs4 import BeautifulSoup

    response = requests.get(
        url,
        timeout=30,
        headers={"User-Agent": "MRPG/1.0 (+https://github.com/AhmdArFa/TuT-AI)"},
    )
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    title = soup.title.string.strip() if soup.title and soup.title.string else url

    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "form"]):
        tag.decompose()
    root = soup.find("article") or soup.find("main") or soup.body or soup
    blocks = []
    for element in root.find_all(["h1", "h2", "h3", "p", "li"]):
        line = _clean(element.get_text(" ", strip=True))
        if len(line) >= 40 or element.name.startswith("h"):
            if line:
                blocks.append(line)
    return Resource(source=url, kind="web", title=_clean(title), blocks=blocks)


def load_resource(source: str) -> Resource:
    """Load one resource from a local path or an http(s) URL."""
    if source.startswith(("http://", "https://")):
        return _from_web(source)

    path = Path(source).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"Resource not found: {source}")
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _from_pdf(path)
    if ext == ".docx":
        return _from_docx(path)
    if ext == ".pptx":
        return _from_pptx(path)
    if ext in {".txt", ".md", ".markdown"}:
        return _from_text(path)
    raise ValueError(
        f"Unsupported resource type '{ext}'. Supported: {sorted(SUPPORTED_EXTENSIONS)} and URLs."
    )


def load_resources(sources: list[str]) -> list[Resource]:
    """Load many resources, skipping empties but failing loudly on bad paths."""
    resources = []
    for source in sources:
        resource = load_resource(source)
        if resource.text:
            resources.append(resource)
    if not resources:
        raise ValueError("None of the provided resources produced any text.")
    return resources
