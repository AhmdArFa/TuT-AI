"""End-to-end smoke test: ingest -> outline (extractive) -> PPTX deck.

Runs with no network and no GPU, so it passes anywhere:

    python tests/smoke_test.py
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from mrpg.deck import render_deck
from mrpg.ingest import load_resources
from mrpg.outline import build_outline_extractive
from pptx import Presentation


def main() -> int:
    root = Path(__file__).resolve().parent.parent
    sources = [
        str(root / "examples" / "resources" / "sovereign_ai_notes.md"),
        str(root / "examples" / "resources" / "document_intelligence.txt"),
    ]

    resources = load_resources(sources)
    assert len(resources) == 2, f"expected 2 resources, got {len(resources)}"
    assert all(r.text for r in resources), "a resource produced no text"
    print(f"ingest OK: {[r.origin_label for r in resources]}")

    outline = build_outline_extractive(
        resources, "Sovereign Document Intelligence", slides=6
    )
    assert outline.slides, "outline has no slides"
    assert all(s.bullets for s in outline.slides), "a slide has no bullets"
    print(f"outline OK: {len(outline.slides)} slides")

    out = root / "examples" / "demo_output.pptx"
    render_deck(outline, out)
    deck = Presentation(str(out))
    expected = len(outline.slides) + 2  # title + closing
    actual = len(list(deck.slides))
    assert actual == expected, f"expected {expected} slides, got {actual}"
    print(f"deck OK: {out} ({expected} slides total)")
    print("SMOKE TEST PASSED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
